"""Construction d'un export de materiaux (Excel + un dossier par materiau)
a partir d'un dossier "plat" de fichiers .BM3MAT.

Dossier source attendu (un ou plusieurs fichiers par materiau) :
    <ID>.BM3MAT                   -> materiau HQ (copie aussi en MQ)
    <ID>_MQ.BM3MAT / _LQ          -> niveaux de detail explicites (optionnels)
    <ID>.jpg / .jpeg / .png       -> miniature du materiau (optionnelle)

Sortie (meme structure que "fichier bm3mat/exemple") :
    <out_dir>/export-products.xlsx            (Products, Assets, Utilities)
    <out_dir>/<ID>/HQ-Lod_0_std.BM3MAT        copie du fichier source
    <out_dir>/<ID>/MQ-Lod_1_std.BM3MAT        copie du HQ (sauf _MQ fourni)
    <out_dir>/<ID>/LQ-Lod_2_std.BM3MAT        HQ aux textures reduites (sauf _LQ fourni)
    <out_dir>/<ID>/thumbnail-512.jpg          texture albedo (ou couleur unie)

Un .BM3MAT est un zip { manifest.json, binary.bin } : le manifeste decrit
les materiaux et leurs images, stockees bout a bout dans binary.bin
(byteOffset/byteLength, alignes sur 4 octets, hash = SHA-256 en 8 entiers
32 bits big-endian). Les exports de reference montrent que :
  - HQ et MQ sont le meme fichier ;
  - LQ reprend le meme manifeste, chaque texture reduite a 256 px ; la carte
    de normales y passe en PNG (sans perte), l'albedo reste en JPG ;
  - la miniature est la texture albedo en 512 px, ou une couleur unie pour
    un materiau sans texture.
Les fichiers source sont copies, jamais deplaces ni modifies.
"""

import hashlib
import io
import json
import shutil
import struct
import zipfile
from pathlib import Path

import openpyxl
import wx

from . import utilities_data
from .bm3_builder import ASSETS_HEADER, IMAGE_SUFFIXES, LODS, THUMBNAIL_STEM, _LOD_SUFFIX

MAT_SUFFIX = ".bm3mat"
MAT_XLSX_NAME = "export-products.xlsx"

LOD_FILE_NAMES = {"HQ": "HQ-Lod_0_std.BM3MAT", "MQ": "MQ-Lod_1_std.BM3MAT", "LQ": "LQ-Lod_2_std.BM3MAT"}
LQ_TEXTURE_SIZE = 256
THUMBNAIL_SIZE = 512
JPEG_QUALITY = 90

CATEGORY_ID = "Material"
DEFAULT_PRODUCT_TYPE = "632+Textile"

PRODUCTS_HEADER = ["Product ID", "Reference", "Product Type", "Brand", "Category ID", "Name (en-GB)"]

_MANIFEST = "manifest.json"


class SourceMaterial:
    def __init__(self, product_id):
        self.product_id = product_id
        self.files = {}  # {"HQ"|"MQ"|"LQ": Path}
        self.thumbnail = None  # Path ou None
        self.texture_count = 0
        # Saisis par l'utilisateur avant l'ecriture de l'Excel.
        self.reference = ""
        self.product_type = DEFAULT_PRODUCT_TYPE
        self.brand = ""
        self.category = CATEGORY_ID
        self.name = ""

    @property
    def source_lod(self):
        """Fichier de reference (le plus detaille disponible) : copie en HQ,
        base du LQ et de la miniature."""
        return next((lod for lod in LODS if lod in self.files), None)


# --------------------------------------------------------------- .BM3MAT --

def _read_bm3mat(path):
    """(manifeste, binary.bin, infos zip) d'un .BM3MAT."""
    with zipfile.ZipFile(path) as z:
        manifest = json.loads(z.read(_MANIFEST).decode("utf-8"))
        uris = [b.get("uri") for b in manifest.get("binaries", [])]
        binary = z.read(uris[0]) if uris else b""
        infos = {i.filename: i for i in z.infolist()}
        others = {name: z.read(name) for name in infos if name != _MANIFEST and name not in uris}
    return manifest, binary, infos, others


def _image_hash(data):
    return list(struct.unpack(">8I", hashlib.sha256(data).digest()))


def _image_bytes(manifest, binary, image):
    if image.get("binary", 0) != 0 or len(manifest.get("binaries", [])) > 1:
        raise ValueError("multi-binary .BM3MAT files are not supported")
    return binary[image["byteOffset"]:image["byteOffset"] + image["byteLength"]]


def _load_image(data):
    image = wx.Image(io.BytesIO(data))
    if not image.IsOk():
        raise ValueError("unreadable texture image")
    return image


def _encode(image, fmt):
    out = io.BytesIO()
    if fmt == "jpg":
        image.SetOption(wx.IMAGE_OPTION_QUALITY, JPEG_QUALITY)
        ok = image.SaveFile(out, wx.BITMAP_TYPE_JPEG)
    else:
        ok = image.SaveFile(out, wx.BITMAP_TYPE_PNG)
    if not ok:
        raise ValueError(f"unable to encode a {fmt} texture")
    return out.getvalue()


def _fit(image, size):
    """Image reduite pour que son plus grand cote fasse `size` px (jamais
    agrandie)."""
    w, h = image.GetWidth(), image.GetHeight()
    scale = size / max(w, h)
    if scale >= 1:
        return image
    return image.Scale(max(1, round(w * scale)), max(1, round(h * scale)), wx.IMAGE_QUALITY_HIGH)


def _normal_images(manifest):
    """Indices des images utilisees comme carte de normales."""
    textures = manifest.get("textures", [])
    found = set()
    for material in manifest.get("materials", []):
        tex = (material.get("normal") or {}).get("texture")
        if tex is not None and tex < len(textures):
            found.add(textures[tex]["image"])
    return found


def _albedo(manifest):
    """(index d'image ou None, couleur lineaire [r, g, b]) de l'albedo du
    premier materiau."""
    materials = manifest.get("materials", [])
    albedo = materials[0].get("albedo", {}) if materials else {}
    colour = albedo.get("value", [1, 1, 1])
    tex = albedo.get("texture")
    textures = manifest.get("textures", [])
    image = textures[tex]["image"] if tex is not None and tex < len(textures) else None
    return image, colour


def build_lq(src, dst):
    """Ecrit dans `dst` la version LQ du .BM3MAT `src` : meme manifeste,
    chaque texture reduite a LQ_TEXTURE_SIZE px. Sans texture a reduire, le
    fichier est copie tel quel (comme dans les exports de reference)."""
    manifest, binary, infos, others = _read_bm3mat(src)
    images = manifest.get("images", [])
    if not images:
        shutil.copy2(src, dst)
        return

    normals = _normal_images(manifest)
    new_binary = bytearray()
    for index, image in enumerate(images):
        picture = _load_image(_image_bytes(manifest, binary, image))
        fmt = "png" if index in normals or picture.HasAlpha() or picture.HasMask() else "jpg"
        data = _encode(_fit(picture, LQ_TEXTURE_SIZE), fmt)
        new_binary.extend(b"\0" * (-len(new_binary) % 4))
        image.update(binary=0, byteOffset=len(new_binary), byteLength=len(data), format=fmt, hash=_image_hash(data))
        new_binary.extend(data)
    manifest["binaries"][0]["byteLength"] = len(new_binary)

    bin_uri = manifest["binaries"][0]["uri"]
    entries = {
        _MANIFEST: json.dumps(manifest, separators=(",", ":"), ensure_ascii=False).encode("utf-8"),
        bin_uri: bytes(new_binary),
        **others,
    }
    with zipfile.ZipFile(dst, "w", zipfile.ZIP_DEFLATED) as z:
        for name, data in entries.items():
            info = zipfile.ZipInfo(name, infos[name].date_time if name in infos else (1980, 1, 1, 0, 0, 0))
            info.compress_type = zipfile.ZIP_DEFLATED
            z.writestr(info, data)


def _linear_to_srgb(value):
    value = min(max(float(value), 0.0), 1.0)
    srgb = value * 12.92 if value <= 0.0031308 else 1.055 * value ** (1 / 2.4) - 0.055
    return round(srgb * 255)


def render_thumbnail(src, dst):
    """Ecrit dans `dst` (JPG) la miniature du .BM3MAT `src` : sa texture
    albedo reduite a THUMBNAIL_SIZE px, ou une image unie de la couleur
    albedo (convertie de lineaire en sRGB) s'il n'a pas de texture."""
    manifest, binary, _infos, _others = _read_bm3mat(src)
    index, colour = _albedo(manifest)
    if index is not None:
        picture = _fit(_load_image(_image_bytes(manifest, binary, manifest["images"][index])), THUMBNAIL_SIZE)
    else:
        picture = wx.Image(THUMBNAIL_SIZE, THUMBNAIL_SIZE)
        r, g, b = (_linear_to_srgb(c) for c in (list(colour) + [1, 1, 1])[:3])
        picture.SetRGB(wx.Rect(0, 0, THUMBNAIL_SIZE, THUMBNAIL_SIZE), r, g, b)
    if picture.HasAlpha():
        picture.ClearAlpha()
    Path(dst).write_bytes(_encode(picture, "jpg"))


# ----------------------------------------------------------- scan/write --

def scan_material_folder(src_dir, log=None):
    """Liste les .BM3MAT du dossier (non recursif), les regroupe par
    Product ID et lit leur manifeste. Retourne une liste de SourceMaterial
    triee par ID. `log`, si fourni, recoit les avertissements (texte)."""

    src_dir = Path(src_dir)
    if not src_dir.is_dir():
        raise FileNotFoundError(f"Folder not found: {src_dir}")

    materials = {}
    images = []
    for path in sorted(src_dir.iterdir()):
        if not path.is_file():
            continue
        suffix = path.suffix.lower()
        if suffix in IMAGE_SUFFIXES:
            images.append(path)
            continue
        if suffix != MAT_SUFFIX:
            continue
        match = _LOD_SUFFIX.match(path.stem)
        product_id, lod = (match["id"], match["lod"].upper()) if match else (path.stem, "HQ")
        material = materials.setdefault(product_id, SourceMaterial(product_id))
        if lod in material.files and log:
            log(f"{product_id}: several {lod} files, keeping {material.files[lod].name} (ignored: {path.name}).")
        material.files.setdefault(lod, path)

    unused_images = []
    for image in images:
        material = materials.get(image.stem)
        if material and not material.thumbnail:
            material.thumbnail = image
        else:
            unused_images.append(image.name)
    if unused_images and log:
        log(
            f"Image(s) not matched to any material, ignored: {', '.join(unused_images)} "
            "(name a thumbnail <ProductID>.jpg to attach it)."
        )

    for material in materials.values():
        lod = material.source_lod
        if lod != "HQ" and log:
            log(f"{material.product_id}: no HQ file, {lod} used as the reference file.")
        manifest = _read_bm3mat(material.files[lod])[0]
        material.texture_count = len(manifest.get("images", []))

    return [materials[pid] for pid in sorted(materials)]


def write_material_export(materials, out_dir, log=None):
    """Ecrit un dossier par materiau (HQ/MQ/LQ + miniature) dans out_dir et
    <out_dir>/export-products.xlsx. Retourne le chemin de l'Excel. out_dir
    doit etre distinct du dossier source des .BM3MAT."""

    out_dir = Path(out_dir)
    for material in materials:
        for path in list(material.files.values()) + [material.thumbnail]:
            if path is not None and path.parent.resolve() == out_dir.resolve():
                raise ValueError("The output folder must be different from the .BM3MAT files folder.")

    out_dir.mkdir(parents=True, exist_ok=True)

    wb = openpyxl.Workbook()
    ws_products = wb.active
    ws_products.title = "Products"
    ws_products.append(PRODUCTS_HEADER)
    ws_assets = wb.create_sheet("Assets")
    ws_assets.append(ASSETS_HEADER)

    for material in materials:
        pid = material.product_id
        product_dir = out_dir / pid
        product_dir.mkdir(exist_ok=True)
        reference = material.files[material.source_lod]

        done = []
        for lod in LODS:
            dst = product_dir / LOD_FILE_NAMES[lod]
            if lod in material.files:
                shutil.copy2(material.files[lod], dst)
            elif lod == "LQ":
                build_lq(reference, dst)
                done.append("LQ built")
            else:
                shutil.copy2(reference, dst)
        if material.thumbnail:
            name = THUMBNAIL_STEM + material.thumbnail.suffix.lower()
            shutil.copy2(material.thumbnail, product_dir / name)
        else:
            name = THUMBNAIL_STEM + ".jpg"
            render_thumbnail(reference, product_dir / name)
            done.append("thumbnail rendered")
        if log:
            log(f"{pid}: HQ/MQ/LQ{''.join(', ' + d for d in done)} -> {product_dir}")

        ws_products.append([
            pid, material.reference or "", material.product_type or "", material.brand or "",
            material.category or "", material.name or "",
        ])
        ws_assets.append([pid] + [f"{pid}/{LOD_FILE_NAMES[lod]}" for lod in LODS] + ["", f"{pid}/{name}", None])

    ws_utilities = wb.create_sheet("Utilities")
    ws_utilities.append(utilities_data.HEADER)
    for row in utilities_data.ROWS:
        ws_utilities.append(row)

    for ws in (ws_products, ws_assets):
        ws.column_dimensions["A"].width = 20
    for col in ("B", "C", "D", "E", "F"):
        ws_products.column_dimensions[col].width = 25
    ws_utilities.column_dimensions["A"].width = 30

    out_xlsx = out_dir / MAT_XLSX_NAME
    wb.save(out_xlsx)
    return out_xlsx
