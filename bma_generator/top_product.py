"""Construction d'un "top produit" : le produit catalogue qui regroupe les
assemblages .BMA d'une gamme (ex. tous les modules d'un canape).

Contrairement aux autres exports, il n'a ni fichier 3D ni feuille Assets :
un seul Excel (Products, Parameters, Utilities), comme
test_top_produits/top_Product.xlsx. Sa feuille Parameters porte :
  - "modules" (obligatoire) : les assemblages proposes (IDs <ID>_ass de
    l'export bma), Default = l'un d'eux, Values = tous, joints par "+" ;
  - chaque parametre materiau des modules (ex. structureColor), avec les
    materiaux choisis (voir material_values.py) ou "null" sinon.
"""

from pathlib import Path

import openpyxl

from . import material_values as mv
from . import utilities_data

TOP_XLSX_NAME = "top_Product.xlsx"
MODULES_PARAM = "modules"
CATEGORY_ID = "Furniture"

# Ordre des colonnes de l'exemple de reference (Category ID avant Brand).
PRODUCTS_HEADER = ["Product ID", "Reference", "Product Type", "Category ID", "Brand"]
PARAMETERS_HEADER = [
    "Delete", "Product ID", "Parameter Definition ID", "Parameter ID", "Type",
    "Parameter Translation Key", "Editable", "Default", "Values", "Allow any value",
    "Value Translation Key", "Magnitude", "Tags", "Free Tags", "Step", "Index UI",
    "In or Out", "Visibility", "Force",
]

# Parametres administratifs des assemblages, jamais repris sur le top.
_ADMIN_PARAMS = {"module", "package"}


class BmaModule:
    def __init__(self, assembly_id, reference):
        self.assembly_id = assembly_id
        self.reference = reference
        self.material_params = []  # parametres materiau (ex. ["structureColor"])


def read_bma_export(path):
    """Assemblages d'un export bma (dossier ou export-products-bma.xlsx) :
    liste de BmaModule dans l'ordre de la feuille Products, avec les
    parametres materiau (type Product hors module/package) de chacun."""
    path = Path(path)
    if path.is_dir():
        candidates = sorted(p for p in path.glob("*.xlsx") if not p.name.startswith("~$"))
        if not candidates:
            raise FileNotFoundError(f"No .xlsx file found in {path}")
        path = next((p for p in candidates if p.name == "export-products-bma.xlsx"), candidates[0])

    wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
    try:
        modules = {}
        rows = wb["Products"].iter_rows(values_only=True)
        header = [str(h).strip() if h is not None else "" for h in next(rows, [])]
        col = {name: i for i, name in enumerate(header)}
        if "Product ID" not in col:
            raise ValueError(f"No 'Product ID' column in the Products sheet of {path}")
        for row in rows:
            pid = row[col["Product ID"]] if col["Product ID"] < len(row) else None
            if not pid:
                continue
            ref_i = col.get("Reference")
            reference = row[ref_i] if ref_i is not None and ref_i < len(row) else None
            modules[str(pid)] = BmaModule(str(pid), "" if reference is None else str(reference))

        if "Parameters" in wb.sheetnames:
            rows = wb["Parameters"].iter_rows(values_only=True)
            header = [str(h).strip() if h is not None else "" for h in next(rows, [])]
            col = {name: i for i, name in enumerate(header)}
            for row in rows:
                module = modules.get(str(row[col["Product ID"]]))
                param, ptype = row[col["Parameter ID"]], row[col["Type"]]
                if module and param and ptype == "Product" and param not in _ADMIN_PARAMS:
                    if param not in module.material_params:
                        module.material_params.append(param)
    finally:
        wb.close()
    return list(modules.values())


def write_top_product(out_xlsx, product_id, reference, product_type, brand, modules,
                      material_params=(), material_values=None, category=CATEGORY_ID):
    """Ecrit l'Excel du top produit. `modules` : (default_id, [ids]) des
    assemblages proposes ; `material_params` : parametres materiau a exposer
    (ex. ["structureColor"]), remplis depuis `material_values` ou a "null".
    Retourne le chemin ecrit."""
    default_module, module_ids = modules
    if not product_id:
        raise ValueError("The top product needs a Product ID.")
    if not module_ids:
        raise ValueError("Choose at least one module (.BMA assembly).")

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Products"
    ws.append(PRODUCTS_HEADER)
    ws.append([product_id, reference or "", product_type or "", category or "", brand or ""])

    ws_params = wb.create_sheet("Parameters")
    ws_params.append(PARAMETERS_HEADER)
    ws_params.append([
        None, product_id, None, MODULES_PARAM, "Product", "", "None",
        default_module if default_module in module_ids else module_ids[0], mv.SEPARATOR.join(module_ids),
        "", "", "", "", "", None, None, "in", None, None,
    ])
    for name in material_params:
        chosen = mv.cells(material_values, name)
        default, values, allow_any = (*chosen, mv.ALLOW_ANY_VALUE) if chosen else ("null", "null", "true")
        ws_params.append([
            None, product_id, None, name, "Product", name, "Read & Write",
            default, values, allow_any, None, None, None, None, None, None, "in", False, None,
        ])

    ws_utilities = wb.create_sheet("Utilities")
    ws_utilities.append(utilities_data.HEADER)
    for row in utilities_data.ROWS:
        ws_utilities.append(row)

    ws.column_dimensions["A"].width = 30
    for c in ("B", "C", "D", "E"):
        ws.column_dimensions[c].width = 22
    ws_params.column_dimensions["B"].width = 30
    ws_utilities.column_dimensions["A"].width = 30

    out_xlsx = Path(out_xlsx)
    out_xlsx.parent.mkdir(parents=True, exist_ok=True)
    wb.save(out_xlsx)
    return out_xlsx
