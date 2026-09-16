"""Generation des fichiers .BMA et de l'Excel de sortie a partir des produits
bm3 et de leur role d'ancrage."""

import json
import shutil
import uuid
from pathlib import Path

import openpyxl

from . import templates as tpl

PACKAGE_SUFFIX = "_p"
ASSEMBLY_SUFFIX = "_ass"

PARAMETERS_HEADER = [
    "Delete", "Product ID", "Parameter Definition ID", "Parameter ID", "Type",
    "Parameter Translation Key", "Editable", "Default", "Values",
    "Allow any value", "Value Translation Key", "Magnitude", "Tags",
    "Free Tags", "Step", "Index UI", "In or Out", "Visibility", "Force",
]

ASSETS_HEADER = [
    "Product ID", "Model 3D HQ", "Model 3D MQ", "Model 3D LQ", "Model 2D",
    "Thumbnail", "Media",
]

PRODUCTS_BASE_HEADER = ["Product ID", "Reference", "Product Type", "Brand"]

PRODUCT_PARAM_TYPE = "Product"


def build_products_header(bm3_products):
    """Header de la feuille Products de sortie : les colonnes de base, plus
    toute colonne supplementaire presente dans le bm3 source (Name,
    Commercial Description, etc.), dans leur ordre d'origine - pour que le
    bma ait exactement les memes colonnes que le bm3."""

    header = list(PRODUCTS_BASE_HEADER)
    seen = set(header)
    for p in bm3_products:
        for name in p.source_header:
            if name not in seen:
                header.append(name)
                seen.add(name)
    return header


def new_uuid():
    return str(uuid.uuid4()).upper()


def _output_row(assembly_id, param_id, ptype, translation_key, editable, default, values,
                 allow_any_value, value_translation_key, magnitude, tags, free_tags, step,
                 index_ui, in_or_out, visibility, force):
    """Construit une ligne de la feuille Parameters de sortie. Pour les
    parametres de type 'Product' (references a un autre produit/materiau),
    Default/Values sont mis a null et Allow any value a 'true' : ces valeurs
    sont resolues dynamiquement par le configurateur, pas figees ici."""

    if ptype == PRODUCT_PARAM_TYPE:
        # Litteralement le mot "null" (pas une cellule vide) : ces valeurs
        # sont resolues dynamiquement par le configurateur, pas figees ici.
        default, values, allow_any_value = "null", "null", "true"

    return [
        None, assembly_id, None, param_id, ptype, translation_key, editable,
        default, values, allow_any_value, value_translation_key, magnitude,
        tags, free_tags, step, index_ui, in_or_out, visibility, force,
    ]


def build_bma_json(product, role, extra_parameters=None):
    """Construit le JSON .BMA pour ce produit/role.

    Seuls les parametres bm3 de type "Product" (references a un autre
    produit/materiau - la couleur, quel que soit son nom, structureColor/
    sofaColor/...) sont redeclares dans "parameters" (exposes sans valeur,
    overridables) et overloades sur MonModule ET Package (meme logique que
    dans les fichiers de reference valides). Les dimensions numeriques du
    bm3 (depth/width/height...) ne sont PAS redeclarees ici : un exemple
    .bma valide (role central) n'en a aucune - les relations y accedent
    directement via componentDependencies (MonModule.width, etc.), sans
    passer par un parametre expose au niveau de l'assemblage.

    `extra_parameters`, lui, reste gere normalement (ex: le
    profondeurModuleAttache de la chaise longue, un parametre technique
    propre a l'assemblage, pas un parametre du bm3 source).

    Retourne (data, package_id, warnings) : warnings est toujours vide ici,
    conserve pour compatibilite avec les appelants existants.
    """
    relations, anchors = tpl.build_role_definition(role)
    anchors = [{"uuid": new_uuid(), **a} for a in anchors]

    package_id = f"{product.product_id}{PACKAGE_SUFFIX}"
    warnings = []

    parameters = [
        {
            "type": "component",
            "name": "monModule",
            "value": {
                "protocol": "product",
                "referenceValue": {"dbId": product.product_id},
            },
        },
    ]
    module_overloads = []
    package_overloads = []

    for param in product.parameters:
        if param.get("type") != PRODUCT_PARAM_TYPE:
            continue
        pid = param["id"]
        parameters.append({"type": "component", "name": pid})
        module_overloads.append({"parameter": pid, "type": "component", "value": pid})
        package_overloads.append({"parameter": pid, "type": "component", "value": pid})

    if extra_parameters:
        parameters.extend(extra_parameters)

    parameters.append(
        {
            "type": "component",
            "name": "package",
            "value": {
                "protocol": "product",
                "referenceValue": {"dbId": package_id},
            },
        }
    )

    data = {
        "version": 7,
        "uuid": new_uuid(),
        "name": product.product_id,
        "parameters": parameters,
        "relations": relations,
        "components": [
            {
                "name": "MonModule",
                "position": {"x": 0, "y": 0, "z": 0},
                "directionY": {"x": 0, "y": 1, "z": 0},
                "directionZ": {"x": 0, "y": 0, "z": 1},
                "overloads": module_overloads,
                "reference": "monModule",
                "positioningMode": 0,
                "rotationAngle": 0,
            },
            {
                "name": "Package",
                "position": {"x": 0, "y": 0, "z": 0},
                "directionY": {"x": 0, "y": 1, "z": 0},
                "directionZ": {"x": 0, "y": 0, "z": 1},
                "overloads": package_overloads,
                "reference": "package",
                "positioningMode": 0,
                "rotationAngle": 0,
            },
        ],
        "anchors": anchors,
    }
    return data, package_id, warnings


def _normalize_role_info(value):
    """Accepte soit une chaine (role seul), soit un dict {'role':..,
    'ref_depth':..}, et renvoie toujours (role, ref_depth)."""
    if isinstance(value, dict):
        return value.get("role"), value.get("ref_depth")
    return value, None


def find_reference_depth(product, role, roles_by_id, products_by_id):
    """Deduit la profondeur des modules qui vont se connecter a ce produit
    (role chaise longue), en cherchant parmi les autres produits du lot ceux
    dont le role expose un tag que ce produit peut recevoir."""

    wanted_tags = tpl.receive_tags(role)
    depths = []
    for other_id, other_role in roles_by_id.items():
        if other_id == product.product_id or not other_role or other_role not in tpl.ALL_ROLES:
            continue
        if other_role == tpl.CHAISE_LONGUE:
            # Une autre chaise longue n'est pas un module "standard" : on ne
            # veut que la profondeur des modules normaux auxquels elle se
            # connecte, jamais celle d'une chaise longue voisine.
            continue
        if tpl.exposed_tags(other_role) & wanted_tags:
            other_product = products_by_id.get(other_id)
            if other_product is not None and other_product.depth is not None:
                depths.append(other_product.depth)

    if not depths:
        return None, []

    rounded = {round(d, 1) for d in depths}
    warnings = []
    if len(rounded) > 1:
        warnings.append(
            f"{product.product_id} (chaise longue): inconsistent depths among compatible "
            f"modules ({sorted(rounded)}), using {depths[0]}."
        )
    return depths[0], warnings


def generate(bm3_products, roles, bm3_dir, out_dir, brand_override=None):
    """Genere l'arborescence fichier bma/ pour tous les produits ayant un role.

    `roles` : dict {product_id: role} ou {product_id: {"role":.., "ref_depth":..}}.
    `brand_override` : si fourni, remplace la colonne Brand du bm3 source
    (prioritaire) ; sinon la valeur du bm3 est conservee.

    Retourne (generated, skipped, rows_products, rows_assets, rows_parameters,
    warnings, products_header).
    """
    bm3_dir = Path(bm3_dir)
    out_dir = Path(out_dir)
    out_dir.mkdir(parents=True, exist_ok=True)

    products_header = build_products_header(bm3_products)
    extra_columns = products_header[len(PRODUCTS_BASE_HEADER):]

    products_by_id = {p.product_id: p for p in bm3_products}
    roles_by_id = {}
    ref_depth_overrides = {}
    for pid, value in roles.items():
        role, ref_depth = _normalize_role_info(value)
        roles_by_id[pid] = role
        if ref_depth is not None:
            ref_depth_overrides[pid] = ref_depth

    generated = []
    skipped = []
    warnings = []
    rows_products = []
    rows_assets = []
    rows_parameters = []

    for product in bm3_products:
        role = roles_by_id.get(product.product_id)
        if not role or role not in tpl.ALL_ROLES:
            skipped.append(product.product_id)
            continue

        extra_parameters = None
        extra_output_rows = []
        assembly_id = f"{product.product_id}{ASSEMBLY_SUFFIX}"

        if role == tpl.CHAISE_LONGUE:
            ref_depth = ref_depth_overrides.get(product.product_id)
            if ref_depth is None:
                ref_depth, dep_warnings = find_reference_depth(product, role, roles_by_id, products_by_id)
                warnings.extend(dep_warnings)
            if ref_depth is None:
                warnings.append(
                    f"{product.product_id} (chaise longue): no compatible module found in the batch "
                    "to deduce the reference depth; set it manually in the roles file. "
                    "Value used by default: the product's own depth (no offset)."
                )
                ref_depth = product.depth
            extra_parameters = [
                {
                    "type": "number",
                    "name": tpl.REFERENCE_DEPTH_PARAM_NAME,
                    "value": ref_depth,
                    "magnitude": 0,
                }
            ]
            extra_output_rows.append(
                _output_row(
                    assembly_id, tpl.REFERENCE_DEPTH_PARAM_NAME, "Real (continuous)",
                    tpl.REFERENCE_DEPTH_PARAM_NAME, "None", ref_depth, ref_depth,
                    "false", "", "Length", "", "", None, None, "in", False, None,
                )
            )

        data, package_id, param_warnings = build_bma_json(product, role, extra_parameters)
        warnings.extend(param_warnings)

        product_dir = out_dir / assembly_id
        product_dir.mkdir(parents=True, exist_ok=True)

        content = json.dumps(data, indent=1, ensure_ascii=False)
        for lod in ("HQ", "MQ", "LQ"):
            (product_dir / f"{lod}-root.BMA").write_text(content, encoding="utf-8")

        src_thumb = bm3_dir / product.asset_id / "thumbnail-512.jpg"
        if src_thumb.exists():
            shutil.copyfile(src_thumb, product_dir / "thumbnail-512.jpg")

        reference_numeric = product.reference
        if reference_numeric and "-" in reference_numeric:
            reference_numeric = reference_numeric.split("-", 1)[1]
        if reference_numeric and reference_numeric.endswith("_bm3"):
            reference_numeric = reference_numeric[: -len("_bm3")]

        # Un Brand saisi dans l'interface remplace explicitement celui du bm3
        # source (l'utilisateur veut pouvoir forcer la marque a l'export).
        brand = brand_override or product.brand or ""
        extra_values = [product.source_columns.get(name) for name in extra_columns]
        rows_products.append([assembly_id, reference_numeric, product.product_type, brand] + extra_values)
        rows_assets.append([
            assembly_id,
            f"{assembly_id}/HQ-root.BMA",
            f"{assembly_id}/MQ-root.BMA",
            f"{assembly_id}/LQ-root.BMA",
            "",
            f"{assembly_id}/thumbnail-512.jpg",
            None,
        ])

        rows_parameters.append(
            _output_row(
                assembly_id, "module", PRODUCT_PARAM_TYPE, "module", "None",
                product.product_id, product.product_id, "false", "", "", "", "",
                None, None, "in", False, None,
            )
        )
        # Meme filtre que build_bma_json : seuls les parametres de type
        # "Product" sont exposes sur l'assemblage genere, pour que l'Excel
        # reste coherent avec ce que le .BMA declare reellement.
        for param in product.parameters:
            if param.get("type") != PRODUCT_PARAM_TYPE:
                continue
            rows_parameters.append(
                _output_row(
                    assembly_id, param["id"], param.get("type"), param.get("translation_key"),
                    param.get("editable"), param.get("default"), param.get("values"),
                    param.get("allow_any_value"), param.get("value_translation_key"),
                    param.get("magnitude"), param.get("tags"), param.get("free_tags"),
                    param.get("step"), param.get("index_ui"), param.get("in_or_out"),
                    param.get("visibility"), param.get("force"),
                )
            )
        rows_parameters.extend(extra_output_rows)
        rows_parameters.append(
            _output_row(
                assembly_id, "package", PRODUCT_PARAM_TYPE, "package", "None",
                package_id, package_id, "false", "", "", "", "",
                None, None, "in", None, None,
            )
        )

        generated.append(product.product_id)

    return generated, skipped, rows_products, rows_assets, rows_parameters, warnings, products_header


def write_output_excel(rows_products, rows_assets, rows_parameters, utilities_source_xlsx, out_xlsx,
                        products_header=None):
    wb = openpyxl.Workbook()

    ws = wb.active
    ws.title = "Products"
    ws.append(products_header or PRODUCTS_BASE_HEADER)
    for row in rows_products:
        ws.append(row)

    ws = wb.create_sheet("Assets")
    ws.append(ASSETS_HEADER)
    for row in rows_assets:
        ws.append(row)

    ws = wb.create_sheet("Parameters")
    ws.append(PARAMETERS_HEADER)
    for row in rows_parameters:
        ws.append(row)

    ws = wb.create_sheet("Utilities")
    src_wb = openpyxl.load_workbook(utilities_source_xlsx, data_only=True)
    src_ws = src_wb["Utilities"]
    for row in src_ws.iter_rows(values_only=True):
        ws.append(row)

    wb.save(out_xlsx)
